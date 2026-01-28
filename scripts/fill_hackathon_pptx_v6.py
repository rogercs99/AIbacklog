from pathlib import Path
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.text import PP_ALIGN

HOME = Path.home()
TEMPLATE = HOME / "Downloads/aXetHackathon_project_template (2).pptx"
OUTPUT = HOME / "Downloads/Req2Backlog_AI_Hackathon_v6.pptx"

TITLE_FONT = "Calibri"
BODY_FONT = "Calibri"
BODY_SIZE = Pt(18)
TITLE_SIZE = Pt(34)
LINE_SPACING = 1.2
SPACE_AFTER = Pt(6)

prs = Presentation(str(TEMPLATE))

# Helper to apply paragraph styling

def style_paragraph(p, size=BODY_SIZE, bold=False, align=PP_ALIGN.LEFT):
    p.alignment = align
    p.line_spacing = LINE_SPACING
    p.space_after = SPACE_AFTER
    for run in p.runs:
        run.font.name = BODY_FONT
        run.font.size = size
        run.font.bold = bold

# Slide 0 — Title
slide0 = prs.slides[0]
if slide0.shapes[0].has_text_frame:
    tf = slide0.shapes[0].text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "Req2Backlog AI"
    run.font.name = TITLE_FONT
    run.font.size = TITLE_SIZE
    run.font.bold = True
    p.alignment = PP_ALIGN.CENTER

    p2 = tf.add_paragraph()
    run2 = p2.add_run()
    run2.text = "From requirements to an executable backlog with traceability"
    run2.font.name = BODY_FONT
    run2.font.size = Pt(20)
    p2.alignment = PP_ALIGN.CENTER
    p2.line_spacing = 1.2

# Slide 1 — Team
slide1 = prs.slides[1]
if slide1.shapes[1].has_text_frame:
    tf = slide1.shapes[1].text_frame
    tf.clear()
    lines = [
        "Participant: Roger Campos i Sans",
        "Role: Individual Participant (Solution Designer & Developer)",
        "Region / Function: NTT DATA – DT DAR Spain – aXet GenAI",
        "Challenge Category: 1A – AI Productivity Booster",
        "Project Name: Req2Backlog AI",
    ]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = line
        style_paragraph(p)

# Slide 2 — Challenge
slide2 = prs.slides[2]
if slide2.shapes[1].has_text_frame:
    tf = slide2.shapes[1].text_frame
    tf.clear()
    lines = [
        "Requirements live in long docs, emails, and meeting notes.",
        "Turning them into a structured backlog is slow and error‑prone.",
        "Teams lose time aligning scope, acceptance criteria, and dependencies.",
        "Need: a faster, clearer path from requirements to execution.",
    ]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = line
        style_paragraph(p)

# Slide 3 — Project
slide3 = prs.slides[3]
if slide3.shapes[1].has_text_frame:
    tf = slide3.shapes[1].text_frame
    tf.clear()
    lines = [
        "Tool that makes individual work faster, easier, and smarter.",
        "Generates Epics → Stories → Tasks with acceptance criteria + estimates.",
        "Adds traceability to document chunks and highlights missing info.",
        "Reusable & scalable: modular flows + Jira/Rally exports.",
        "Tech: aXet.flows (GenAI), Next.js UI, SQLite persistence.",
    ]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = line
        style_paragraph(p)

# Slide 4 — Demo title stays, just style if needed
# Slide 5 — Demo
slide5 = prs.slides[5]
if slide5.shapes[1].has_text_frame:
    tf = slide5.shapes[1].text_frame
    tf.clear()
    lines = [
        "Upload a requirements document (PDF/DOCX/TXT).",
        "Click Generate backlog.",
        "Review summary, assumptions, Epics/Stories/Tasks.",
        "Export to Jira/Rally CSV.",
        "Demo video: 3–5 minutes.",
    ]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = line
        style_paragraph(p)

# Slide 6 — Outcomes
slide6 = prs.slides[6]
if slide6.shapes[1].has_text_frame:
    tf = slide6.shapes[1].text_frame
    tf.clear()
    lines = [
        "Backlog creation time reduced from hours to minutes.",
        "Improved consistency and traceability for scope decisions.",
        "Faster alignment with stakeholders and delivery teams.",
        "Next: domain templates, reconciliation UX, more exports.",
    ]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = line
        style_paragraph(p)

prs.save(str(OUTPUT))
print(f"Saved: {OUTPUT}")
